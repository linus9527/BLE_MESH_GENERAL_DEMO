from __future__ import annotations

from datetime import date
from pathlib import Path

from PIL import Image as PILImage

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    HRFlowable,
    Image as RLImage,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "output"
PDF_PATH = OUT_DIR / "四合一Demo保护盒尺寸说明.pdf"
DOCX_PATH = OUT_DIR / "四合一Demo保护盒尺寸说明.docx"

# 四张物料原始尺寸图
IMG_1 = Path(r"C:\Users\19542\AppData\Local\Temp\codex-clipboard-fe405249-51d2-4d66-818d-5d1d69832dea.jpg")
IMG_2 = Path(r"C:\Users\19542\AppData\Local\Temp\codex-clipboard-4ab19121-d46c-466d-8e74-5037334dc80b.jpg")
IMG_3 = Path(r"C:\Users\19542\AppData\Local\Temp\codex-clipboard-0a0ed248-48bb-4312-a41a-f0d314841e89.jpg")
IMG_4 = Path(r"C:\Users\19542\AppData\Local\Temp\codex-clipboard-59c0eadc-9622-4faf-ac10-87cfcab43252.jpg")

# 三张组合 Demo 实测图
IMG_DEMO_1 = Path(r"C:\Users\19542\AppData\Local\Temp\codex-clipboard-56f15f9d-0d9a-4ef7-9cd3-648ca0ca1b1c.jpg")
IMG_DEMO_2 = Path(r"C:\Users\19542\AppData\Local\Temp\codex-clipboard-550bc185-0439-4b76-86da-ad5730501a20.jpg")
IMG_DEMO_3 = Path(r"C:\Users\19542\AppData\Local\Temp\codex-clipboard-0a0a596a-6b5d-4d32-ab5d-c0ec9e84d841.jpg")

ALL_IMAGES = [IMG_1, IMG_2, IMG_3, IMG_4, IMG_DEMO_1, IMG_DEMO_2, IMG_DEMO_3]

NAVY = colors.HexColor("#17324D")
BLUE = colors.HexColor("#1F6F8B")
LIGHT_BLUE = colors.HexColor("#EAF3F7")
PALE = colors.HexColor("#F5F7F9")
GRAY = colors.HexColor("#64748B")
GRID = colors.HexColor("#CBD5E1")


def validate_inputs() -> None:
    missing = [str(path) for path in ALL_IMAGES if not path.exists()]
    if missing:
        raise FileNotFoundError("缺少图片：\n" + "\n".join(missing))


def register_fonts() -> tuple[str, str]:
    pairs = [
        (r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\msyhbd.ttc"),
        (r"C:\Windows\Fonts\simhei.ttf", r"C:\Windows\Fonts\simhei.ttf"),
        (r"C:\Windows\Fonts\simsun.ttc", r"C:\Windows\Fonts\simsun.ttc"),
    ]
    for regular, bold in pairs:
        if Path(regular).exists() and Path(bold).exists():
            try:
                pdfmetrics.registerFont(TTFont("CJK", regular))
                pdfmetrics.registerFont(TTFont("CJK-Bold", bold))
                return "CJK", "CJK-Bold"
            except Exception:
                pass
    raise RuntimeError("未找到可用的中文字体")


def rl_image(path: Path, max_w: float, max_h: float) -> RLImage:
    with PILImage.open(path) as img:
        width, height = img.size
    scale = min(max_w / width, max_h / height)
    return RLImage(str(path), width=width * scale, height=height * scale)


def styles(font: str, bold_font: str) -> dict[str, ParagraphStyle]:
    sample = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "TitleCN", parent=sample["Title"], fontName=bold_font, fontSize=23,
            leading=31, textColor=NAVY, alignment=TA_LEFT, spaceAfter=4 * mm,
        ),
        "sub": ParagraphStyle(
            "SubCN", parent=sample["BodyText"], fontName=font, fontSize=9,
            leading=14, textColor=GRAY, spaceAfter=4 * mm,
        ),
        "h1": ParagraphStyle(
            "H1CN", parent=sample["Heading1"], fontName=bold_font, fontSize=15,
            leading=21, textColor=NAVY, spaceBefore=2 * mm, spaceAfter=3 * mm,
        ),
        "body": ParagraphStyle(
            "BodyCN", parent=sample["BodyText"], fontName=font, fontSize=9.5,
            leading=15.5, textColor=colors.HexColor("#263746"), spaceAfter=2 * mm,
            wordWrap="CJK",
        ),
        "table": ParagraphStyle(
            "TableCN", parent=sample["BodyText"], fontName=font, fontSize=8.3,
            leading=12.5, textColor=colors.HexColor("#263746"), wordWrap="CJK",
        ),
        "head": ParagraphStyle(
            "HeadCN", parent=sample["BodyText"], fontName=bold_font, fontSize=8.3,
            leading=12, textColor=colors.white, alignment=TA_CENTER, wordWrap="CJK",
        ),
        "dimension": ParagraphStyle(
            "DimensionCN", parent=sample["BodyText"], fontName=bold_font, fontSize=27,
            leading=34, textColor=NAVY, alignment=TA_CENTER,
        ),
        "center": ParagraphStyle(
            "CenterCN", parent=sample["BodyText"], fontName=font, fontSize=9,
            leading=13, textColor=GRAY, alignment=TA_CENTER,
        ),
        "caption": ParagraphStyle(
            "CaptionCN", parent=sample["BodyText"], fontName=font, fontSize=7.8,
            leading=11.5, textColor=GRAY, alignment=TA_CENTER, wordWrap="CJK",
        ),
    }


def pdf_header_footer(canvas, doc, font: str, bold_font: str) -> None:
    canvas.saveState()
    page_w, page_h = A4
    canvas.setStrokeColor(colors.HexColor("#DDE5EA"))
    canvas.setLineWidth(0.5)
    canvas.line(doc.leftMargin, page_h - 12 * mm, page_w - doc.rightMargin, page_h - 12 * mm)
    canvas.setFont(bold_font, 7.5)
    canvas.setFillColor(NAVY)
    canvas.drawString(doc.leftMargin, page_h - 9 * mm, "四合一 Demo 保护盒尺寸说明")
    canvas.line(doc.leftMargin, 12 * mm, page_w - doc.rightMargin, 12 * mm)
    canvas.setFont(font, 7.2)
    canvas.setFillColor(GRAY)
    canvas.drawString(doc.leftMargin, 8 * mm, "尺寸单位：mm｜尺寸顺序：长 × 宽 × 高")
    canvas.drawRightString(page_w - doc.rightMargin, 8 * mm, f"第 {doc.page} 页")
    canvas.restoreState()


def reportlab_table(data, col_widths, s, center_cols=()) -> Table:
    rows = [[Paragraph(str(cell), s["head"] if r == 0 else s["table"]) for cell in row] for r, row in enumerate(data)]
    table = Table(rows, colWidths=col_widths, repeatRows=1, hAlign="CENTER")
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), NAVY),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, PALE]),
        ("GRID", (0, 0), (-1, -1), 0.45, GRID),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2.2 * mm),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2.2 * mm),
        ("TOPPADDING", (0, 0), (-1, -1), 2.3 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.3 * mm),
    ]
    for col in center_cols:
        commands.append(("ALIGN", (col, 0), (col, -1), "CENTER"))
    table.setStyle(TableStyle(commands))
    return table


def build_pdf() -> None:
    font, bold_font = register_fonts()
    s = styles(font, bold_font)
    doc = SimpleDocTemplate(
        str(PDF_PATH), pagesize=A4,
        leftMargin=24 * mm, rightMargin=24 * mm,
        topMargin=18 * mm, bottomMargin=17 * mm,
        title="四合一Demo保护盒尺寸说明",
        author="BLE_MESH_BUTTON",
    )
    story = []

    # 第 1 页：只写尺寸结论和用户的取舍。
    story.append(Spacer(1, 8 * mm))
    story.append(Paragraph("四合一开发 Demo", s["sub"]))
    story.append(Paragraph("保护盒尺寸说明", s["title"]))
    story.append(Paragraph(f"供采购查看　｜　{date.today().isoformat()}", s["sub"]))
    size_card = Table(
        [
            [Paragraph("最终决定", ParagraphStyle("Label", parent=s["body"], fontName=bold_font, textColor=BLUE, alignment=TA_CENTER))],
            [Paragraph("90 × 80 × 90 mm", s["dimension"])],
            [Paragraph("长 × 宽 × 高", s["center"])],
        ],
        colWidths=[162 * mm],
    )
    size_card.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), LIGHT_BLUE),
        ("BOX", (0, 0), (-1, -1), 1.1, BLUE),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("TOPPADDING", (0, 0), (-1, -1), 3 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3 * mm),
    ]))
    story.append(size_card)
    story.append(Spacer(1, 7 * mm))
    story.append(Paragraph("一、组合 Demo 尺寸依据", s["h1"]))
    demo_size_data = [
        ["方向", "实物估算", "本次决定", "说明"],
        ["长", "约 80", "90", "在原估算基础上增加约 10 mm。"],
        ["宽", "约 80", "80", "曾考虑改为 90–100 mm，最终仍采用 80 mm。"],
        ["高", "按 90 考虑", "90", "高度方向稍作预留，作为装配余量。"],
    ]
    story.append(reportlab_table(demo_size_data, [24 * mm, 34 * mm, 32 * mm, 72 * mm], s, center_cols=(0, 1, 2)))
    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph("二、宽度考虑与最终取舍", s["h1"]))
    story.append(Paragraph(
        "传感器接线部分会略微超出组合 Demo 的主体轮廓，因此曾考虑把宽度从 80 mm 增加到 90–100 mm，给接线留出更多空间。",
        s["body"],
    ))
    story.append(Paragraph(
        "结合当前实物接线状态，少量超出问题不大。最终决定采用 <b>90 × 80 × 90 mm（长 × 宽 × 高）</b>，不再把宽度增加到 90–100 mm。",
        s["body"],
    ))

    # 第 2 页：三张实测照片及尺寸表。
    story.append(PageBreak())
    story.append(Paragraph("三、组合 Demo 实物测量图片", s["h1"]))
    demo_items = [
        (IMG_DEMO_1, "实物图 1｜宽度约 80 mm"),
        (IMG_DEMO_2, "实物图 2｜高度按 90 mm 考虑"),
        (IMG_DEMO_3, "实物图 3｜长度约 80 mm"),
    ]
    image_row = [rl_image(path, 49 * mm, 99 * mm) for path, _ in demo_items]
    caption_row = [Paragraph(caption, s["caption"]) for _, caption in demo_items]
    photo_table = Table([image_row, caption_row], colWidths=[54 * mm] * 3, hAlign="CENTER")
    photo_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), PALE),
        ("GRID", (0, 0), (-1, -1), 0.45, GRID),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2 * mm),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2 * mm),
        ("TOPPADDING", (0, 0), (-1, -1), 2 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2 * mm),
    ]))
    story.append(photo_table)
    story.append(Spacer(1, 7 * mm))
    story.append(Paragraph("尺寸汇总", s["h1"]))
    summary_data = [
        ["项目", "长度", "宽度", "高度", "单位"],
        ["组合 Demo 实物估算", "约 80", "约 80", "按 90 考虑", "mm"],
        ["保护盒最终决定", "90", "80", "90", "mm"],
    ]
    story.append(reportlab_table(summary_data, [54 * mm, 27 * mm, 27 * mm, 31 * mm, 23 * mm], s, center_cols=(1, 2, 3, 4)))

    # 第 3 页：四张物料的原始尺寸表。
    story.append(PageBreak())
    story.append(Paragraph("四、四项物料原始尺寸表", s["h1"]))
    story.append(Paragraph("以下尺寸均来自所附原始尺寸图片，并统一使用毫米。", s["body"]))
    material_data = [
        ["图号", "物料", "原始尺寸（mm）", "原图说明"],
        ["1", "DC/USB 供电转换板", "71.7 × 48", "安装孔中心距约 65.5 × 42；高度未标注。"],
        ["2", "12 V 电池组（SRJ-1223）", "70 × 55 × 19", "原图标注为 19 × 55 × 70，此处按长 × 宽 × 高排列。"],
        ["3", "Seeed Studio XIAO nRF52840", "22 × 18", "高度未标注。"],
        ["4", "XIAO 转接/驱动板", "42 × 25 × 14", "原图标注长、宽、高。"],
    ]
    story.append(reportlab_table(material_data, [14 * mm, 49 * mm, 40 * mm, 59 * mm], s, center_cols=(0, 2)))

    # 第 4—5 页：四张原始尺寸图放在文档最后。
    story.append(PageBreak())
    story.append(Paragraph("附录｜物料原始尺寸图片（1/2）", s["h1"]))
    story.append(Paragraph("图 1　DC/USB 供电转换板", s["caption"]))
    story.append(Spacer(1, 1.5 * mm))
    story.append(rl_image(IMG_1, 162 * mm, 104 * mm))
    story.append(Spacer(1, 4 * mm))
    story.append(HRFlowable(width="100%", thickness=0.5, color=GRID))
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph("图 2　12 V 电池组 SRJ-1223", s["caption"]))
    story.append(Spacer(1, 1.5 * mm))
    story.append(rl_image(IMG_2, 162 * mm, 76 * mm))

    story.append(PageBreak())
    story.append(Paragraph("附录｜物料原始尺寸图片（2/2）", s["h1"]))
    story.append(Paragraph("图 3　Seeed Studio XIAO nRF52840", s["caption"]))
    story.append(Spacer(1, 1.5 * mm))
    story.append(rl_image(IMG_3, 155 * mm, 101 * mm))
    story.append(Spacer(1, 4 * mm))
    story.append(HRFlowable(width="100%", thickness=0.5, color=GRID))
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph("图 4　XIAO 转接/驱动板", s["caption"]))
    story.append(Spacer(1, 1.5 * mm))
    story.append(rl_image(IMG_4, 155 * mm, 96 * mm))

    def callback(canvas, active_doc):
        pdf_header_footer(canvas, active_doc, font, bold_font)

    doc.build(story, onFirstPage=callback, onLaterPages=callback)


def set_run(run, size=9.5, bold=False, color="263746") -> None:
    run.font.name = "微软雅黑"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_run(paragraph, text, size=9.5, bold=False, color="263746"):
    run = paragraph.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    return run


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margin(cell, value=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def clear(cell) -> None:
    cell.text = ""
    cell.paragraphs[0].text = ""


def add_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(5)
    add_run(p, text, size=15, bold=True, color="17324D")


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.25
    add_run(p, text, size=9.5)


def add_picture(paragraph, path: Path, max_width_mm: float, max_height_mm: float) -> None:
    with PILImage.open(path) as img:
        width, height = img.size
    target_width = max_width_mm
    if target_width * height / width > max_height_mm:
        target_width = max_height_mm * width / height
    paragraph.add_run().add_picture(str(path), width=Mm(target_width))


def style_word_table(table, header=True) -> None:
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell_margin(cell)
            if header and row_idx == 0:
                shade(cell, "17324D")
            elif row_idx % 2 == 0:
                shade(cell, "F5F7F9")


def fill_word_table(table, data, center_cols=()) -> None:
    for row_idx, row in enumerate(data):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            clear(cell)
            p = cell.paragraphs[0]
            if col_idx in center_cols or row_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_run(p, value, size=8.2, bold=(row_idx == 0), color="FFFFFF" if row_idx == 0 else "263746")
    style_word_table(table)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(paragraph, "第 ", size=8, color="64748B")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    add_run(paragraph, " 页", size=8, color="64748B")


def build_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(17)
    section.bottom_margin = Mm(17)
    section.left_margin = Mm(24)
    section.right_margin = Mm(24)
    section.header_distance = Mm(6)
    section.footer_distance = Mm(6)

    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(9.5)

    header = section.header.paragraphs[0]
    add_run(header, "四合一 Demo 保护盒尺寸说明", size=8, bold=True, color="17324D")
    add_page_field(section.footer.paragraphs[0])

    # 第 1 页
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    add_run(p, "四合一开发 Demo", size=10, color="64748B")
    p = doc.add_paragraph()
    add_run(p, "保护盒尺寸说明", size=23, bold=True, color="17324D")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    add_run(p, f"供采购查看　｜　{date.today().isoformat()}", size=9, color="64748B")

    card = doc.add_table(rows=3, cols=1)
    card.style = "Table Grid"
    card.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell in [row.cells[0] for row in card.rows]:
        shade(cell, "EAF3F7")
        cell_margin(cell, 120)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(card.cell(0, 0).paragraphs[0], "最终决定", size=10, bold=True, color="1F6F8B")
    add_run(card.cell(1, 0).paragraphs[0], "90 × 80 × 90 mm", size=26, bold=True, color="17324D")
    add_run(card.cell(2, 0).paragraphs[0], "长 × 宽 × 高", size=9, color="64748B")

    add_heading(doc, "一、组合 Demo 尺寸依据")
    demo_size_data = [
        ["方向", "实物估算", "本次决定", "说明"],
        ["长", "约 80", "90", "在原估算基础上增加约 10 mm。"],
        ["宽", "约 80", "80", "曾考虑改为 90–100 mm，最终仍采用 80 mm。"],
        ["高", "按 90 考虑", "90", "高度方向稍作预留，作为装配余量。"],
    ]
    table = doc.add_table(rows=len(demo_size_data), cols=4)
    fill_word_table(table, demo_size_data, center_cols=(0, 1, 2))

    add_heading(doc, "二、宽度考虑与最终取舍")
    add_body(doc, "传感器接线部分会略微超出组合 Demo 的主体轮廓，因此曾考虑把宽度从 80 mm 增加到 90–100 mm，给接线留出更多空间。")
    add_body(doc, "结合当前实物接线状态，少量超出问题不大。最终决定采用 90 × 80 × 90 mm（长 × 宽 × 高），不再把宽度增加到 90–100 mm。")

    # 第 2 页
    doc.add_page_break()
    add_heading(doc, "三、组合 Demo 实物测量图片")
    photos = doc.add_table(rows=2, cols=3)
    photos.style = "Table Grid"
    photos.alignment = WD_TABLE_ALIGNMENT.CENTER
    items = [
        (IMG_DEMO_1, "实物图 1｜宽度约 80 mm"),
        (IMG_DEMO_2, "实物图 2｜高度按 90 mm 考虑"),
        (IMG_DEMO_3, "实物图 3｜长度约 80 mm"),
    ]
    for col, (path, caption) in enumerate(items):
        for row in range(2):
            cell = photos.cell(row, col)
            clear(cell)
            shade(cell, "F5F7F9")
            cell_margin(cell, 70)
        p = photos.cell(0, col).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_picture(p, path, 47, 96)
        p = photos.cell(1, col).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p, caption, size=7.5, color="64748B")

    add_heading(doc, "尺寸汇总")
    summary_data = [
        ["项目", "长度", "宽度", "高度", "单位"],
        ["组合 Demo 实物估算", "约 80", "约 80", "按 90 考虑", "mm"],
        ["保护盒最终决定", "90", "80", "90", "mm"],
    ]
    table = doc.add_table(rows=len(summary_data), cols=5)
    fill_word_table(table, summary_data, center_cols=(1, 2, 3, 4))

    # 第 3 页
    doc.add_page_break()
    add_heading(doc, "四、四项物料原始尺寸表")
    add_body(doc, "以下尺寸均来自所附原始尺寸图片，并统一使用毫米。")
    material_data = [
        ["图号", "物料", "原始尺寸（mm）", "原图说明"],
        ["1", "DC/USB 供电转换板", "71.7 × 48", "安装孔中心距约 65.5 × 42；高度未标注。"],
        ["2", "12 V 电池组（SRJ-1223）", "70 × 55 × 19", "原图为 19 × 55 × 70，此处按长 × 宽 × 高排列。"],
        ["3", "Seeed Studio XIAO nRF52840", "22 × 18", "高度未标注。"],
        ["4", "XIAO 转接/驱动板", "42 × 25 × 14", "原图标注长、宽、高。"],
    ]
    table = doc.add_table(rows=len(material_data), cols=4)
    fill_word_table(table, material_data, center_cols=(0, 2))

    # 最后两页只放四张物料原始尺寸图。
    doc.add_page_break()
    add_heading(doc, "附录｜物料原始尺寸图片（1/2）")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "图 1　DC/USB 供电转换板", size=8, color="64748B")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_picture(p, IMG_1, 160, 102)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "图 2　12 V 电池组 SRJ-1223", size=8, color="64748B")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_picture(p, IMG_2, 160, 72)

    doc.add_page_break()
    add_heading(doc, "附录｜物料原始尺寸图片（2/2）")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "图 3　Seeed Studio XIAO nRF52840", size=8, color="64748B")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_picture(p, IMG_3, 153, 99)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "图 4　XIAO 转接/驱动板", size=8, color="64748B")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_picture(p, IMG_4, 153, 92)

    doc.core_properties.title = "四合一Demo保护盒尺寸说明"
    doc.core_properties.author = "BLE_MESH_BUTTON"
    doc.save(DOCX_PATH)


def main() -> None:
    validate_inputs()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_pdf()
    build_docx()
    print(PDF_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
